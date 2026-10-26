from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


PARSER_VERSION = "x1-parser-0.1.0"
MAX_CHARS = 1200
OVERLAP_CHARS = 120


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class SourceSegment:
    text: str
    locator: dict[str, Any]
    section_title: str = ""


@dataclass(frozen=True)
class ParsedChunk:
    sequence_no: int
    text: str
    locator: dict[str, Any]
    section_title: str
    char_start: int
    char_end: int
    token_estimate: int
    checksum: str


@dataclass(frozen=True)
class ParsedPayload:
    parser_name: str
    title: str
    page_count: int
    segments: list[SourceSegment]
    chunks: list[ParsedChunk]
    warnings: list[dict[str, str]] = field(default_factory=list)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(char for char in text if char in "\n\t" or unicodedata.category(char) != "Cc")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(path: Path) -> tuple[str, str]:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("text_encoding_unsupported", "TXT文件应使用UTF-8或GB18030编码")


def _parse_pdf(path: Path) -> tuple[list[SourceSegment], int, list[dict[str, str]]]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentParseError("encrypted_document", "暂不接收加密PDF")
    segments: list[SourceSegment] = []
    warnings: list[dict[str, str]] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            segments.append(SourceSegment(text=text, locator={"kind": "page", "page": page_index}))
        else:
            warnings.append({"code": "page_without_text", "message": f"第{page_index}页未提取到文本"})
    if not segments:
        raise DocumentParseError("no_extractable_text", "PDF没有可提取文本，扫描件需先完成OCR")
    return segments, len(reader.pages), warnings


def _parse_docx(path: Path) -> tuple[list[SourceSegment], int, list[dict[str, str]]]:
    document = Document(str(path))
    segments: list[SourceSegment] = []
    current_heading = ""
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.lower().startswith("heading") or style_name.startswith("标题"):
            current_heading = text[:255]
        segments.append(SourceSegment(
            text=text,
            locator={"kind": "paragraph", "paragraph": index},
            section_title=current_heading,
        ))
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                text = normalize_text(cell.text)
                if text:
                    segments.append(SourceSegment(
                        text=text,
                        locator={
                            "kind": "table_cell",
                            "table": table_index,
                            "row": row_index,
                            "cell": cell_index,
                        },
                        section_title=current_heading,
                    ))
    if not segments:
        raise DocumentParseError("no_extractable_text", "DOCX没有可提取文本")
    return segments, 1, []


def _parse_pptx(path: Path) -> tuple[list[SourceSegment], int, list[dict[str, str]]]:
    presentation = Presentation(str(path))
    segments: list[SourceSegment] = []
    warnings: list[dict[str, str]] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = normalize_text(shape.text)
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = normalize_text(cell.text)
                        if text:
                            texts.append(text)
        joined = normalize_text("\n".join(texts))
        if joined:
            title = normalize_text(slide.shapes.title.text)[:255] if slide.shapes.title else ""
            segments.append(SourceSegment(
                text=joined,
                locator={"kind": "slide", "slide": slide_index},
                section_title=title,
            ))
        else:
            warnings.append({"code": "slide_without_text", "message": f"第{slide_index}页未提取到文本"})
    if not segments:
        raise DocumentParseError("no_extractable_text", "PPTX没有可提取文本")
    return segments, len(presentation.slides), warnings


def _parse_txt(path: Path) -> tuple[list[SourceSegment], int, list[dict[str, str]]]:
    decoded, encoding = _decode_text(path)
    normalized = normalize_text(decoded)
    if not normalized:
        raise DocumentParseError("no_extractable_text", "TXT没有可提取文本")
    segments: list[SourceSegment] = []
    cursor = 1
    for block in re.split(r"\n\s*\n", normalized):
        text = normalize_text(block)
        if not text:
            continue
        line_count = text.count("\n") + 1
        segments.append(SourceSegment(
            text=text,
            locator={"kind": "line_range", "line_start": cursor, "line_end": cursor + line_count - 1},
        ))
        cursor += line_count + 1
    return segments, 1, [{"code": "text_encoding", "message": f"按{encoding}解析"}]


def chunk_segments(segments: list[SourceSegment]) -> list[ParsedChunk]:
    chunks: list[ParsedChunk] = []
    sequence = 1
    step = MAX_CHARS - OVERLAP_CHARS
    for segment in segments:
        start = 0
        while start < len(segment.text):
            hard_end = min(len(segment.text), start + MAX_CHARS)
            end = hard_end
            if hard_end < len(segment.text):
                boundary = max(
                    segment.text.rfind("。", start + MAX_CHARS // 2, hard_end),
                    segment.text.rfind("\n", start + MAX_CHARS // 2, hard_end),
                    segment.text.rfind("；", start + MAX_CHARS // 2, hard_end),
                )
                if boundary > start:
                    end = boundary + 1
            text = segment.text[start:end].strip()
            if text:
                chunks.append(ParsedChunk(
                    sequence_no=sequence,
                    text=text,
                    locator={**segment.locator, "segment_char_start": start, "segment_char_end": end},
                    section_title=segment.section_title,
                    char_start=start,
                    char_end=end,
                    token_estimate=max(1, math.ceil(len(text) / 2.2)),
                    checksum=sha256(text.encode("utf-8")).hexdigest(),
                ))
                sequence += 1
            if end >= len(segment.text):
                break
            start = max(start + step, end - OVERLAP_CHARS)
    return chunks


def parse_document(path: Path, extension: str, original_name: str) -> ParsedPayload:
    parser_name = extension.removeprefix(".")
    if extension == ".pdf":
        segments, page_count, warnings = _parse_pdf(path)
    elif extension == ".docx":
        segments, page_count, warnings = _parse_docx(path)
    elif extension == ".pptx":
        segments, page_count, warnings = _parse_pptx(path)
    elif extension == ".txt":
        segments, page_count, warnings = _parse_txt(path)
    else:
        raise DocumentParseError("unsupported_extension", "没有可用的文档解析器")
    chunks = chunk_segments(segments)
    return ParsedPayload(
        parser_name=parser_name,
        title=Path(original_name).stem[:255],
        page_count=page_count,
        segments=segments,
        chunks=chunks,
        warnings=warnings,
    )
