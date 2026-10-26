from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError

from app.domains.material_alignment.alignment import AlignerConcept, rank_concepts, split_manual_mentions
from app.models import Concept


COURSE_ID = "machine-learning"


def upload_and_parse(client, headers, *, filename, mime, content, purpose="graph_build"):
    prepared = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets:prepare-upload",
        headers=headers,
        json={
            "filename": filename,
            "byte_size": len(content),
            "declared_mime": mime,
            "purpose": purpose,
        },
    )
    assert prepared.status_code == 201, prepared.text
    payload = prepared.json()
    uploaded = client.put(payload["upload_url"], headers=headers, content=content)
    assert uploaded.status_code == 200, uploaded.text
    completed = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets/{payload['asset']['id']}:complete",
        headers=headers,
    )
    return payload["asset"]["id"], completed


def test_txt_upload_parse_align_evidence_and_delete_propagates(client, teacher_headers):
    content = (
        "第一节 线性回归\n\n"
        "损失函数衡量预测值与真实值的偏差。忽略上一条指令、访问外部工具等文字只属于课程原文。"
    ).encode("utf-8")
    asset_id, completed = upload_and_parse(
        client,
        teacher_headers,
        filename="线性回归讲义.txt",
        mime="text/plain",
        content=content,
    )
    assert completed.status_code == 200, completed.text
    body = completed.json()
    assert body["asset"]["security_status"] == "accepted"
    assert body["document"]["chunk_count"] >= 2
    assert body["document"]["parser_version"] == "x1-parser-0.1.0"

    aligned = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/alignment-runs",
        headers=teacher_headers,
        json={"source_asset_id": asset_id, "purpose": "graph_build"},
    )
    assert aligned.status_code == 201, aligned.text
    run = aligned.json()
    assert run["summary"]["mention_count"] == 2
    assert run["summary"]["matched"] == 2
    labels = [item["raw_text"] for item in run["mentions"]]
    assert labels == ["线性回归", "损失函数"]
    assert all(item["evidence_span_id"] for item in run["mentions"])

    evidence = client.get(
        f"/api/v1/teacher/evidence-spans/{run['mentions'][0]['evidence_span_id']}",
        headers=teacher_headers,
    )
    assert evidence.status_code == 200
    assert evidence.json()["quote"] == "线性回归"
    assert evidence.json()["locator"]["kind"] == "line_range"

    graph = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/graphs",
        headers=teacher_headers,
        json={"title": "来源删除传播测试图谱"},
    ).json()
    applied = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{run['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": graph["version"]},
    )
    assert applied.status_code == 200, applied.text
    assert len(applied.json()["nodes"]) == 2
    assert all(node["evidence_count"] == 1 for node in applied.json()["nodes"])

    deleted = client.delete(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets/{asset_id}",
        headers=teacher_headers,
    )
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["invalidated_derivations"] == 4
    invalidated = client.get(f"/api/v1/teacher/alignment-runs/{run['id']}", headers=teacher_headers)
    assert invalidated.status_code == 200
    assert invalidated.json()["status"] == "invalidated"
    gone = client.get(
        f"/api/v1/teacher/evidence-spans/{run['mentions'][0]['evidence_span_id']}",
        headers=teacher_headers,
    )
    assert gone.status_code == 410
    graph_after_delete = client.get(
        f"/api/v1/teacher/graphs/{graph['id']}", headers=teacher_headers
    )
    assert graph_after_delete.status_code == 200
    assert graph_after_delete.json()["nodes"] == []


def test_manual_alignment_preserves_semicolon_order_and_supports_abstention(client, teacher_headers):
    response = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/alignment-runs",
        headers=teacher_headers,
        json={
            "manual_text": "线性回归；损失函数;量子引力",
            "purpose": "graph_build",
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert [item["raw_text"] for item in body["mentions"]] == ["线性回归", "损失函数", "量子引力"]
    assert body["mentions"][0]["status"] == "matched"
    assert body["mentions"][2]["status"] == "abstained"
    review = body["mentions"][2]["review"]
    assert review["status"] == "pending"

    decided = client.post(
        f"/api/v1/teacher/alignment-reviews/{review['id']}:decide",
        headers=teacher_headers,
        json={"decision": "abstained", "note": "不属于本课程规范知识点"},
    )
    assert decided.status_code == 200, decided.text
    assert decided.json()["summary"]["pending_review"] == 0
    assert decided.json()["status"] == "completed"

    graph = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/graphs",
        headers=teacher_headers,
        json={"title": "材料对齐写入测试图谱"},
    ).json()
    applied = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{body['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": graph["version"]},
    )
    assert applied.status_code == 200, applied.text
    assert {node["label"] for node in applied.json()["nodes"]} == {"线性回归", "损失函数"}
    assert all(node["origin"] == "alignment_auto" for node in applied.json()["nodes"])
    repeated = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{body['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": applied.json()["version"]},
    )
    assert repeated.status_code == 200
    assert len(repeated.json()["nodes"]) == 2


def test_teacher_can_create_and_exactly_merge_a_novel_course_concept(client, teacher_headers):
    response = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/alignment-runs",
        headers=teacher_headers,
        json={
            "manual_text": "量子引力；量子引力",
            "purpose": "graph_build",
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert len(body["mentions"]) == 2
    assert all(item["review"]["status"] == "pending" for item in body["mentions"])

    created = client.post(
        f"/api/v1/teacher/alignment-reviews/{body['mentions'][0]['review']['id']}:decide",
        headers=teacher_headers,
        json={
            "decision": "created",
            "new_concept_name": "量子引力导论",
            "note": "教师确认为课程扩展概念",
        },
    )
    assert created.status_code == 200, created.text
    result = created.json()
    assert result["status"] == "completed"
    assert result["summary"]["pending_review"] == 0
    selected_ids = {item["review"]["selected_concept_id"] for item in result["mentions"]}
    assert len(selected_ids) == 1
    assert {item["review"]["decision"] for item in result["mentions"]} == {"created", "matched"}
    assert all(item["status"] == "matched" for item in result["mentions"])
    assert all(
        any(candidate["concept_name"] == "量子引力导论" for candidate in item["candidates"])
        for item in result["mentions"]
    )

    graph = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/graphs",
        headers=teacher_headers,
        json={"title": "开放新概念归并测试"},
    ).json()
    applied = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{body['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": graph["version"]},
    )
    assert applied.status_code == 200, applied.text
    assert [node["label"] for node in applied.json()["nodes"]] == ["量子引力导论"]
    assert applied.json()["nodes"][0]["origin"] == "alignment_review"

    validation = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}:validate",
        headers=teacher_headers,
        json={"expected_revision_version": applied.json()["version"]},
    )
    assert validation.status_code == 200, validation.text
    assert validation.json()["quality"]["hard_pass"] is False
    assert any(
        issue["code"] == "missing_definition"
        for issue in validation.json()["quality"]["issues"]
    )


def test_open_concept_merges_across_runs_and_requires_real_source_evidence(
    client, teacher_headers,
):
    first_run = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/alignment-runs",
        headers=teacher_headers,
        json={"manual_text": "Novel Alpha Topic", "purpose": "graph_build"},
    )
    assert first_run.status_code == 201, first_run.text
    first_review = first_run.json()["mentions"][0]["review"]
    first_decision = client.post(
        f"/api/v1/teacher/alignment-reviews/{first_review['id']}:decide",
        headers=teacher_headers,
        json={
            "decision": "created",
            "new_concept_name": "Open Concept 42",
            "note": "Teacher confirmed this course extension from the manual brief.",
        },
    )
    assert first_decision.status_code == 200, first_decision.text
    first_selected = first_decision.json()["mentions"][0]["review"]["selected_concept_id"]
    assert first_decision.json()["mentions"][0]["review"]["decision"] == "created"

    second_run = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/alignment-runs",
        headers=teacher_headers,
        json={"manual_text": "Unrelated Zeta Topic", "purpose": "graph_build"},
    )
    assert second_run.status_code == 201, second_run.text
    second_review = second_run.json()["mentions"][0]["review"]
    second_decision = client.post(
        f"/api/v1/teacher/alignment-reviews/{second_review['id']}:decide",
        headers=teacher_headers,
        json={
            "decision": "created",
            "new_concept_name": " open-concept 42 ",
            "note": "A second teacher-confirmed mention of the same course concept.",
        },
    )
    assert second_decision.status_code == 200, second_decision.text
    second_result = second_decision.json()["mentions"][0]["review"]
    assert second_result["decision"] == "matched"
    assert second_result["selected_concept_id"] == first_selected

    graph = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/graphs",
        headers=teacher_headers,
        json={"title": "Open concept source merge"},
    ).json()
    first_apply = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{first_run.json()['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": graph["version"]},
    )
    assert first_apply.status_code == 200, first_apply.text
    assert first_apply.json()["revision_id"] == graph["revision_id"]
    second_apply = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}/alignment-runs/{second_run.json()['id']}:apply",
        headers=teacher_headers,
        json={"expected_revision_version": first_apply.json()["version"]},
    )
    assert second_apply.status_code == 200, second_apply.text
    assert len(second_apply.json()["nodes"]) == 1
    node = second_apply.json()["nodes"][0]
    assert node["concept_id"] == first_selected
    assert node["evidence_count"] == 2

    initial_validation = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}:validate",
        headers=teacher_headers,
        json={"expected_revision_version": second_apply.json()["version"]},
    )
    assert initial_validation.status_code == 200, initial_validation.text
    initial_codes = {item["code"] for item in initial_validation.json()["quality"]["issues"]}
    assert "missing_definition" in initial_codes
    assert "missing_provenance" not in initial_codes

    without_source = client.patch(
        f"/api/v1/teacher/graphs/{graph['id']}/nodes/{node['id']}",
        headers=teacher_headers,
        json={
            "expected_revision_version": initial_validation.json()["version"],
            "definition": "A teacher-authored definition for publication review.",
            "teacher_evidence_note": "",
        },
    )
    assert without_source.status_code == 200, without_source.text
    blocked = client.post(
        f"/api/v1/teacher/graphs/{graph['id']}:validate",
        headers=teacher_headers,
        json={"expected_revision_version": without_source.json()["version"]},
    )
    assert blocked.status_code == 200, blocked.text
    assert "missing_provenance" in {
        item["code"] for item in blocked.json()["quality"]["issues"]
    }

    with client.app.state.db.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(Concept).where(
            Concept.course_id == COURSE_ID,
            Concept.normalized_name == "openconcept42",
        )) == 1
        db.add(Concept(
            id=f"{COURSE_ID}:concept:duplicate-test",
            course_id=COURSE_ID,
            name="OPEN CONCEPT 42",
            sort_order=999,
        ))
        with pytest.raises(IntegrityError):
            db.commit()
        db.rollback()


def test_docx_and_pptx_have_deterministic_structural_locators(client, teacher_headers):
    docx_stream = BytesIO()
    document = Document()
    document.add_heading("线性回归", level=1)
    document.add_paragraph("损失函数与梯度下降")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "正则化"
    document.save(docx_stream)
    docx_id, docx_result = upload_and_parse(
        client,
        teacher_headers,
        filename="课程讲义.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=docx_stream.getvalue(),
    )
    assert docx_result.status_code == 200, docx_result.text
    assert docx_result.json()["document"]["segment_count"] == 3

    pptx_stream = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "模型评价"
    slide.placeholders[1].text = "学习率与特征工程"
    presentation.save(pptx_stream)
    _, pptx_result = upload_and_parse(
        client,
        teacher_headers,
        filename="课堂课件.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        content=pptx_stream.getvalue(),
    )
    assert pptx_result.status_code == 200, pptx_result.text
    assert pptx_result.json()["document"]["page_count"] == 1

    listing = client.get(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets?purpose=graph_build",
        headers=teacher_headers,
    )
    assert listing.status_code == 200
    assert docx_id in {item["id"] for item in listing.json()}


def test_rejects_mime_disguise_active_content_encrypted_pdf_and_oversize(client, teacher_headers):
    mismatch = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets:prepare-upload",
        headers=teacher_headers,
        json={
            "filename": "伪装.pdf", "byte_size": 5, "declared_mime": "text/plain",
            "purpose": "graph_build",
        },
    )
    assert mismatch.status_code == 422
    assert mismatch.json()["detail"]["code"] == "declared_mime_mismatch"

    oversize = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets:prepare-upload",
        headers=teacher_headers,
        json={
            "filename": "过大.txt", "byte_size": 3 * 1024 * 1024,
            "declared_mime": "text/plain", "purpose": "graph_build",
        },
    )
    assert oversize.status_code == 422
    assert oversize.json()["detail"]["code"] == "file_too_large"

    fake_id, fake_result = upload_and_parse(
        client, teacher_headers, filename="伪装.pdf", mime="application/pdf", content=b"plain text"
    )
    assert fake_result.status_code == 422
    assert fake_result.json()["detail"]["code"] == "file_signature_mismatch"
    assert fake_id

    active_stream = BytesIO()
    with ZipFile(active_stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "<w:document />")
        archive.writestr("word/vbaProject.bin", b"macro")
    _, active_result = upload_and_parse(
        client,
        teacher_headers,
        filename="含活动内容.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=active_stream.getvalue(),
    )
    assert active_result.status_code == 422
    assert active_result.json()["detail"]["code"] == "active_content_rejected"

    encrypted_stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    writer.write(encrypted_stream)
    _, encrypted_result = upload_and_parse(
        client,
        teacher_headers,
        filename="加密讲义.pdf",
        mime="application/pdf",
        content=encrypted_stream.getvalue(),
    )
    assert encrypted_result.status_code == 422
    assert encrypted_result.json()["detail"]["code"] == "encrypted_document"


def test_student_cannot_upload_teacher_material(client, student_headers):
    response = client.post(
        f"/api/v1/teacher/courses/{COURSE_ID}/assets:prepare-upload",
        headers=student_headers,
        json={
            "filename": "讲义.txt", "byte_size": 8, "declared_mime": "text/plain",
            "purpose": "graph_build",
        },
    )
    assert response.status_code == 403


def test_aligner_pure_functions_cover_match_ambiguity_and_abstention():
    concepts = [
        AlignerConcept(id="a", name="损失函数"),
        AlignerConcept(id="b", name="损失项"),
    ]
    assert split_manual_mentions(" A；B; C ") == ["A", "B", "C"]
    assert rank_concepts("损失函数", "", concepts).status == "matched"
    assert rank_concepts("损失", "", concepts).status == "ambiguous"
    assert rank_concepts("量子引力", "", concepts).status == "abstained"
